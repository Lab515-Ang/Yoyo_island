# 產生 docs/美術清單.xlsx 與 docs/美術清單.html（再用 Chrome 印成 PDF）。資料在 art_data.py。
# 用法：python tools/art-list/build_art_list.py
import html
import sys

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(__import__('pathlib').Path(__file__).parent))
import art_data as D  # noqa: E402

FONT = 'Microsoft JhengHei'
DEEP = '4A6A42'
SAGE_TINT = 'E8EEDB'
CREAM = 'FBF6EC'
INPUT = 'FFF4C2'  # 要填的欄（狀態、備註）
thin = Side(style='thin', color='D8D0BC')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD_ROW = 4  # 表頭在第 4 列，資料從第 5 列開始

STATUS_FILL = {'已完成': 'D6EBCB', '製作中': 'FDE7B0', '待驗收': 'FBD9C8', '暫代': 'E6E1D4', '選配': 'EEF0F4', '未開始': 'FFFFFF'}


def sheet(wb, name, title, desc, headers, widths, rows, input_cols=(), status_col=None):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws['A1'] = title
    ws['A1'].font = Font(name=FONT, size=15, bold=True, color=DEEP)
    ws['A2'] = desc
    ws['A2'].font = Font(name=FONT, size=10, color='766B5E')
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    ws['A2'].alignment = Alignment(wrap_text=True, vertical='top')
    ws.row_dimensions[2].height = 32
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=HEAD_ROW, column=c, value=h)
        cell.font = Font(name=FONT, size=10, bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor=DEEP)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[HEAD_ROW].height = 30
    for r, row in enumerate(rows, HEAD_ROW + 1):
        zebra = r % 2 == 0
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = Font(name=FONT, size=10)
            cell.alignment = Alignment(wrap_text=True, vertical='top', horizontal='center' if isinstance(v, (int, float)) else 'left')
            cell.border = BORDER
            if c in input_cols:
                cell.fill = PatternFill('solid', fgColor=INPUT)
            elif zebra:
                cell.fill = PatternFill('solid', fgColor=CREAM)
        if status_col:
            s = row[status_col - 1]
            ws.cell(row=r, column=status_col).fill = PatternFill('solid', fgColor=STATUS_FILL.get(s, INPUT))
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = ws.cell(row=HEAD_ROW + 1, column=3)
    last = HEAD_ROW + max(1, len(rows))
    ws.auto_filter.ref = f'A{HEAD_ROW}:{get_column_letter(len(headers))}{last}'
    if status_col:
        dv = DataValidation(type='list', formula1='"' + ','.join(D.STATUS) + '"', allow_blank=True)
        dv.error, dv.errorTitle = '請從清單選一個狀態', '狀態'
        ws.add_data_validation(dv)
        col = get_column_letter(status_col)
        dv.add(f'{col}{HEAD_ROW + 1}:{col}{HEAD_ROW + 300}')
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


def build():
    animals = D.animals()
    projects = D.projects()
    inv = D.inventory()
    wb = Workbook()
    wb.remove(wb.active)
    ov = wb.create_sheet('總覽')

    # ---- 共通規格 ----
    sheet(wb, '共通規格', '共通規格', '所有美術都要遵守。每一項後面附上原因（多半是實際做過、踩過的破綻）。',
          ['項目', '規格', '說明／原因'], [16, 52, 70], D.GLOBAL)

    # ---- A 動物主圖 ----
    a_rows = [[a['no'], a['id'], a['name'], a['region'], a['shape'], a['stages'], a['stage_names'], a['look'], a['habitat'], a['time'], a['appear'],
               a['level'], a['unlock'], a['batch'], a['priority'], a['file'], a['status'], a['note']] for a in animals]
    A = sheet(wb, 'A 動物主圖', 'A. 動物主圖（每個成長階段 1 張）',
              '正側面、頭朝右；製作時 1024×1024 白底主圖，App 用 512×512 透明（腳底 y＝472，離下緣 40 px）。三個成長階段畫布一樣大：幼年約成年的 60%、亞成年約 80%，腳底同一條線。這張主圖同時是 B 動物動態的唯一像素來源。黃色欄位由製作者更新。',
              ['編號', '代號', '名稱', '區域', '體型', '張數', '成長階段', '外型重點（給畫師／GPT）', '望遠鏡全景', '活動時段', '出現方式', '保育等級', '取得方式', '交圖批次', '優先度', '檔名', '狀態', '備註'],
              [6, 20, 12, 7, 8, 6, 20, 44, 18, 9, 20, 18, 22, 8, 7, 26, 9, 34], a_rows, input_cols=(17, 18), status_col=17)
    A['F4'].comment = Comment('有三個成長階段的動物 3 張；商店、Plus 會員的動物只有一種樣式 1 張。', '棲光小島')
    A['N4'].comment = Comment('1＝新玩家前兩週一定遇到；2＝台灣主線；3＝日本區；4＝商店；5＝Plus 會員。', '棲光小島')

    # ---- B 動物動態 ----
    tpl = {t[1]: t for t in D.TEMPLATES}
    b_rows = []
    for a in animals:
        t = tpl[a['shape_id']]
        flight = '需要（展翅 1 張）' if a['shape_id'] in ('bird', 'raptor') else '—'
        st, note = '未開始', ''
        if a['id'] == 'leopard_cat':
            st, note = '製作中', '成年完成；目前未壓縮 460 KB 要壓到 ≤120 KB；後腿之間一小撮淺色未解；幼年、亞成年未做'
        b_rows.append([a['no'], a['id'], a['name'], a['shape'], a['stages'], t[4], flight, t[6], '≤ 120 KB', a['batch'], a['priority'], st, note])
    sheet(wb, 'B 動物動態', 'B. 動物動態（分層部件＋網格骨架）',
          '規範：docs/規格-動物動態網格.md（不用 Rive／Spine）。每個成長階段各一套部件；同體型共用骨架與動作（見 B-1）。流程：主圖 → 剝層圖（GPT 一次只改一處）→ 切部件 → 網格與權重 → 半尺寸離線檢查 → A8s 實機檢查 → 輸出。',
          ['編號', '代號', '名稱', '體型', '套數（成長階段）', '剝層圖（每套）', '飛行姿勢', '動作', '容量目標（每套）', '交圖批次', '優先度', '狀態', '備註'],
          [6, 20, 12, 8, 10, 40, 16, 50, 12, 8, 7, 9, 44], b_rows, input_cols=(12, 13), status_col=12)
    sheet(wb, 'B-1 體型模板', 'B-1. 體型模板（同體型共用骨架與動作，只換貼圖）',
          '29 隻動物只有 9 種體型，每種體型做 1 套骨架、權重規則與動作。貓科已用石虎驗證；其他體型以第一隻實作確認後再定案。',
          ['體型', '代號', '適用動物', '圖層（部件）', '剝層圖', '骨架', '動作', '狀態'],
          [9, 10, 30, 52, 36, 34, 46, 14], [list(t) for t in D.TEMPLATES], input_cols=(8,))

    # ---- C 望遠鏡全景 ----
    pano_users = {}
    for a in animals:
        pano_users.setdefault((a['habitat'].split('（')[0]), []).append(a['name'])
    c_rows = []
    for i, (code, name, time, batch, pri, st, note) in enumerate(D.PANOS, 1):
        users = '、'.join(x['name'] for x in animals if x['habitat'] == name.split('（')[0] and (x['time'] == time or (time == '白天' and x['time'] == '晨昏')))
        c_rows.append([i, f'pano-{code}-{D.TIME_CODE[time]}.webp', name, time, users or '（之後新增的動物）', '約 3000×1150', batch, pri, st, note])
    sheet(wb, 'C 望遠鏡全景', 'C. 專注畫面的望遠鏡全景（棲地 × 時段）',
          '專注時上方圓形鏡片在橫向全景上左右搜尋、推近拉遠，動物專注一小段時間後在遠方出現、越來越近。全景照「棲地」共用，不照國家分。晨昏先用白天版調色（選配列）。規格見 C-1。',
          ['編號', '檔名', '棲地', '時段', '會用到的動物（晨昏的動物先用白天版）', '尺寸', '交圖批次', '優先度', '狀態', '備註'],
          [6, 34, 24, 8, 52, 13, 8, 7, 9, 34], c_rows, input_cols=(9, 10), status_col=9)
    sheet(wb, 'C-1 全景規格', 'C-1. 望遠鏡全景規格', '構圖與尺寸都和程式的座標對應，改之前先確認。',
          ['項目', '規格', '說明'], [14, 56, 70], D.PANO_SPEC)

    # ---- D 小島場景 ----
    d_rows = [[p['region'], p['order'], p['id'], p['title'], p['points'], p['look'], p['deliver'], p['priority'], p['status'], p['note']] for p in projects]
    sheet(wb, 'D 小島場景', 'D. 我的小島：場景與修復工程',
          '2D 分層插畫：一張底圖＋每個工程一張去背疊圖（同畫布、完全對齊）。投入 25／50／75／100% 時物件依序出現。規格見 D-1。',
          ['區域', '順序', '工程代號', '工程', '需要點數', '完成後的樣子', '要交的圖', '優先度', '狀態', '備註'],
          [16, 6, 16, 20, 10, 36, 40, 7, 9, 36], d_rows, input_cols=(9, 10), status_col=9)
    sheet(wb, 'D-1 場景規格', 'D-1. 小島場景規格', '', ['項目', '規格', '說明'], [14, 56, 70], D.SCENE_SPEC)

    # ---- E 小道具 ----
    sheet(wb, 'E 小道具', 'E. 望遠鏡畫面的小道具', '風格同全景；去背 PNG。',
          ['項目', '代號', '用途', '規格', '優先度', '狀態', '備註'], [16, 16, 48, 26, 7, 9, 30],
          [list(p) for p in D.PROPS], input_cols=(6, 7), status_col=6)

    # ---- F 介面圖示 ----
    sheet(wb, 'F 介面圖示', 'F. 介面圖示', '黏土風格的圖示（金幣、寶石、寶箱）已完成，其餘照同一個風格。線條圖示（選單、按鈕）可以沿用現有的 lucide。',
          ['群組', '項目', '代號', '規格', '出現位置', '優先度', '狀態', '備註'], [10, 46, 18, 26, 26, 7, 9, 40],
          [list(u) for u in D.UI], input_cols=(7, 8), status_col=7)

    # ---- G 插圖與商店素材 ----
    sheet(wb, 'G 插圖與上架', 'G. 插圖、App 圖示與上架素材', '',
          ['類別', '項目', '規格', '優先度', '狀態', '備註'], [12, 44, 46, 7, 9, 50],
          [list(x) for x in D.ILLUS], input_cols=(5, 6), status_col=5)

    # ---- H 地圖 ----
    sheet(wb, 'H 地圖', 'H. 地圖', '', ['項目', '檔案', '尺寸', '優先度', '狀態', '備註'], [30, 38, 12, 7, 9, 40],
          [list(x) for x in D.MAPS], input_cols=(5, 6), status_col=5)

    # ---- 現有素材 ----
    sheet(wb, '現有素材', '現有素材盤點', f'{D.UPDATED} 專案裡已經有的圖（不含程式畫的臨時圖）。',
          ['檔案', '用途', '尺寸', '大小（KB）', '狀態'], [56, 40, 12, 11, 30], [list(x) for x in inv])

    # ---- 交件檢查 ----
    sheet(wb, '交件檢查', '交件檢查清單', '每次交圖前逐項確認。', ['適用', '檢查項目', '確認'], [12, 90, 10],
          [[a, b, ''] for a, b in D.CHECK], input_cols=(3,))

    # ---- 總覽（公式：數量、完成數由各分頁計算） ----
    ov.sheet_view.showGridLines = False
    ov['A1'] = '棲光小島｜美術清單'
    ov['A1'].font = Font(name=FONT, size=18, bold=True, color=DEEP)
    ov['A2'] = f'更新：{D.UPDATED}。P0＝上架前一定要有；P1＝上架前最好有；P2＝上架後再補。黃色欄位（各分頁的「狀態」「備註」）由製作者更新，數量與完成數會自動計算。'
    ov['A2'].font = Font(name=FONT, size=10, color='766B5E')
    ov.merge_cells('A2:H2')
    ov['A2'].alignment = Alignment(wrap_text=True, vertical='top')
    ov.row_dimensions[2].height = 32
    heads = ['類別', '內容', '項目數', '張數／套數', '已完成', '完成率', '最高優先度', '分頁']
    for c, h in enumerate(heads, 1):
        cell = ov.cell(row=HEAD_ROW, column=c, value=h)
        cell.font = Font(name=FONT, size=10, bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor=DEEP)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = BORDER
    n = lambda k: HEAD_ROW + len(k)  # noqa: E731  最後一列
    q = lambda s: f"'{s}'"  # noqa: E731
    lines = [
        ('A. 動物主圖', '29 種動物，每個成長階段 1 張（正側面）', 'A 動物主圖', 'B', 'F', 'Q', 'P0', n(a_rows)),
        ('B. 動物動態', '分層部件＋網格骨架；同體型共用動作', 'B 動物動態', 'B', 'E', 'L', 'P0', n(b_rows)),
        ('C. 望遠鏡全景', '棲地 × 時段（含選配）', 'C 望遠鏡全景', 'B', None, 'I', 'P0', n(c_rows)),
        ('D. 小島場景', '底圖＋修復工程疊圖（台灣 8、日本 6）', 'D 小島場景', 'D', None, 'I', 'P0', n(d_rows)),
        ('E. 小道具', '倒木、腳印、羽毛等', 'E 小道具', 'A', None, 'F', 'P1', n(D.PROPS)),
        ('F. 介面圖示', '貨幣、寶箱、環境音、頭像、選單', 'F 介面圖示', 'B', None, 'G', 'P0', n(D.UI)),
        ('G. 插圖與上架', 'App 圖示、Google Play、引導與攔截插圖', 'G 插圖與上架', 'B', None, 'E', 'P0', n(D.ILLUS)),
        ('H. 地圖', '世界地圖、3D 地圖貼圖', 'H 地圖', 'A', None, 'E', 'P0', n(D.MAPS)),
    ]
    first = HEAD_ROW + 1
    for r, (cat, what, sh, key, qty, stc, pri, last) in enumerate(lines, HEAD_ROW + 1):
        rng = lambda col: f'{q(sh)}!${col}${first}:${col}${last}'  # noqa: E731
        vals = [cat, what, f'=COUNTA({rng(key)})', f'=SUM({rng(qty)})' if qty else f'=COUNTA({rng(key)})',
                f'=COUNTIF({rng(stc)},"已完成")', f'=IF(C{r}=0,0,E{r}/C{r})', pri, sh]
        for c, v in enumerate(vals, 1):
            cell = ov.cell(row=r, column=c, value=v)
            cell.font = Font(name=FONT, size=10, color='008000' if isinstance(v, str) and v.startswith('=') and c in (3, 4, 5) else '000000')
            cell.border = BORDER
            cell.alignment = Alignment(vertical='top', wrap_text=True, horizontal='center' if c >= 3 else 'left')
        ov.cell(row=r, column=6).number_format = '0%'
    tr = HEAD_ROW + len(lines) + 1
    ov.cell(row=tr, column=1, value='合計').font = Font(name=FONT, size=10, bold=True)
    for c, col in ((3, 'C'), (4, 'D'), (5, 'E')):
        cell = ov.cell(row=tr, column=c, value=f'=SUM({col}{first}:{col}{tr - 1})')
        cell.font = Font(name=FONT, size=10, bold=True)
        cell.alignment = Alignment(horizontal='center')
    ov.cell(row=tr, column=6, value=f'=IF(C{tr}=0,0,E{tr}/C{tr})').number_format = '0%'
    for c in range(1, 9):
        ov.cell(row=tr, column=c).border = BORDER
        ov.cell(row=tr, column=c).fill = PatternFill('solid', fgColor=SAGE_TINT)
    ov['D4'].comment = Comment('A：所有成長階段的張數加總；B：部件套數加總（每個成長階段一套）。其他類別等於項目數。', '棲光小島')
    leg = tr + 2
    ov.cell(row=leg, column=1, value='圖例').font = Font(name=FONT, size=10, bold=True, color=DEEP)
    for i, (s, col) in enumerate(STATUS_FILL.items()):
        cell = ov.cell(row=leg + 1 + i, column=1, value=s)
        cell.fill = PatternFill('solid', fgColor=col)
        cell.font = Font(name=FONT, size=10)
        cell.border = BORDER
        ov.cell(row=leg + 1 + i, column=2, value={
            '已完成': '已交圖並接進 App', '製作中': '正在畫或正在接', '待驗收': '已交圖，等實機檢查', '暫代': '目前用程式畫或暫代圖',
            '選配': '有需要再做', '未開始': '還沒開始'}[s]).font = Font(name=FONT, size=10)
    ov.cell(row=leg + 8, column=1, value='綠色數字是公式（自動計算）；黃色欄位請填寫。').font = Font(name=FONT, size=9, color='766B5E')
    for c, w in enumerate([18, 44, 10, 12, 10, 10, 12, 18], 1):
        ov.column_dimensions[get_column_letter(c)].width = w
    wb.calculation.fullCalcOnLoad = True

    out = D.ROOT / 'docs' / '美術清單.xlsx'
    wb.save(out)

    # ---- 總覽的數值（給 PDF；和公式同一個算法） ----
    def done(rows, idx):
        return sum(1 for r in rows if r[idx] == '已完成')
    summary = [
        ('A. 動物主圖', len(a_rows), sum(r[5] for r in a_rows), done(a_rows, 16), 'P0'),
        ('B. 動物動態', len(b_rows), sum(r[4] for r in b_rows), done(b_rows, 11), 'P0'),
        ('C. 望遠鏡全景', len(c_rows), len(c_rows), done(c_rows, 8), 'P0'),
        ('D. 小島場景', len(d_rows), len(d_rows), done(d_rows, 8), 'P0'),
        ('E. 小道具', len(D.PROPS), len(D.PROPS), sum(1 for p in D.PROPS if p[5] == '已完成'), 'P1'),
        ('F. 介面圖示', len(D.UI), len(D.UI), sum(1 for p in D.UI if p[6] == '已完成'), 'P0'),
        ('G. 插圖與上架', len(D.ILLUS), len(D.ILLUS), sum(1 for p in D.ILLUS if p[4] == '已完成'), 'P0'),
        ('H. 地圖', len(D.MAPS), len(D.MAPS), sum(1 for p in D.MAPS if p[4] == '已完成'), 'P0'),
    ]

    # ---- HTML（印成 PDF） ----
    e = lambda v: '' if v is None else html.escape(str(v))  # noqa: E731

    def table(headers, rows, cls=''):
        h = ''.join(f'<th>{e(x)}</th>' for x in headers)
        b = ''
        for r in rows:
            b += '<tr>' + ''.join(f'<td class="{"st st-" + e(v) if v in STATUS_FILL else ""}">{e(v)}</td>' for v in r) + '</tr>'
        return f'<table class="{cls}"><thead><tr>{h}</tr></thead><tbody>{b}</tbody></table>'

    tot = [sum(s[i] for s in summary) for i in (1, 2, 3)]
    parts = [f'<h1>棲光小島｜美術清單</h1><p class="sub">更新：{D.UPDATED}。P0＝上架前一定要有；P1＝上架前最好有；P2＝上架後再補。可編輯版本：docs/美術清單.xlsx。</p>',
             '<h2>總覽</h2>' + table(['類別', '項目數', '張數／套數', '已完成', '最高優先度'], summary + [('合計', *tot, '')], 'ov')]
    parts.append('<h2>共通規格</h2>' + table(['項目', '規格', '說明／原因'], D.GLOBAL))
    parts.append('<h2 class="pb">A. 動物主圖（每個成長階段 1 張）</h2><p class="sub">正側面、頭朝右；製作 1024×1024 白底主圖，App 用 512×512 透明（腳底 y＝472）。幼年約成年的 60%、亞成年約 80%。這張主圖也是動物動態的唯一像素來源。</p>'
                 + table(['#', '代號', '名稱', '區域', '體型', '張數', '成長階段', '外型重點', '全景', '時段', '出現方式', '保育等級', '取得方式', '批次', '優先', '檔名', '狀態', '備註'], a_rows, 'small'))
    parts.append('<h2 class="pb">B. 動物動態（分層部件＋網格骨架）</h2><p class="sub">規範見 docs/規格-動物動態網格.md。每個成長階段一套部件；同體型共用骨架與動作。</p>'
                 + table(['#', '代號', '名稱', '體型', '套數', '剝層圖（每套）', '飛行姿勢', '動作', '容量', '批次', '優先', '狀態', '備註'], b_rows, 'small'))
    parts.append('<h2>B-1. 體型模板</h2>' + table(['體型', '代號', '適用動物', '圖層（部件）', '剝層圖', '骨架', '動作', '狀態'], D.TEMPLATES, 'small'))
    parts.append('<h2 class="pb">C. 望遠鏡全景（棲地 × 時段）</h2>' + table(['#', '檔名', '棲地', '時段', '會用到的動物', '尺寸', '批次', '優先', '狀態', '備註'], c_rows, 'small')
                 + '<h2>C-1. 全景規格</h2>' + table(['項目', '規格', '說明'], D.PANO_SPEC))
    parts.append('<h2 class="pb">D. 我的小島：場景與修復工程</h2>' + table(['區域', '順序', '工程代號', '工程', '點數', '完成後的樣子', '要交的圖', '優先', '狀態', '備註'], d_rows, 'small')
                 + '<h2>D-1. 小島場景規格</h2>' + table(['項目', '規格', '說明'], D.SCENE_SPEC))
    parts.append('<h2 class="pb">E. 望遠鏡畫面的小道具</h2>' + table(['項目', '代號', '用途', '規格', '優先', '狀態', '備註'], D.PROPS, 'small')
                 + '<h2>F. 介面圖示</h2>' + table(['群組', '項目', '代號', '規格', '出現位置', '優先', '狀態', '備註'], D.UI, 'small'))
    parts.append('<h2 class="pb">G. 插圖、App 圖示與上架素材</h2>' + table(['類別', '項目', '規格', '優先', '狀態', '備註'], D.ILLUS, 'small')
                 + '<h2>H. 地圖</h2>' + table(['項目', '檔案', '尺寸', '優先', '狀態', '備註'], D.MAPS, 'small'))
    parts.append('<h2 class="pb">現有素材盤點</h2>' + table(['檔案', '用途', '尺寸', 'KB', '狀態'], inv, 'small')
                 + '<h2>交件檢查清單</h2>' + table(['適用', '檢查項目', '確認'], [(a, b, '☐') for a, b in D.CHECK]))
    css = f"""
@page {{ size: A4 landscape; margin: 12mm 10mm 14mm; }}
body {{ font-family: 'Microsoft JhengHei', 'Noto Sans TC', sans-serif; color: #36302a; font-size: 9.5pt; }}
h1 {{ color: #{DEEP}; font-size: 20pt; margin: 0 0 2mm; }}
h2 {{ color: #{DEEP}; font-size: 13pt; margin: 6mm 0 2mm; border-left: 4px solid #8CA67A; padding-left: 2.5mm; }}
h2.pb {{ break-before: page; margin-top: 0; }}
p.sub {{ color: #766b5e; margin: 0 0 3mm; }}
table {{ width: 100%; border-collapse: collapse; margin-bottom: 3mm; }}
thead {{ display: table-header-group; }}
tr {{ break-inside: avoid; }}
th {{ background: #{DEEP}; color: #fff; font-weight: 700; padding: 1.6mm 1.4mm; text-align: left; border: 1px solid #{DEEP}; }}
td {{ padding: 1.3mm 1.4mm; border: 1px solid #D8D0BC; vertical-align: top; }}
tbody tr:nth-child(even) td {{ background: #{CREAM}; }}
table.small {{ font-size: 7.6pt; }}
table.ov td:nth-child(n+2) {{ text-align: center; }}
table.ov tr:last-child td {{ font-weight: 700; background: #{SAGE_TINT}; }}
""" + ''.join(f'td.st-{s} {{ background: #{c} !important; white-space: nowrap; }}' for s, c in STATUS_FILL.items())
    page = f'<!doctype html><html lang="zh-Hant"><head><meta charset="utf-8"><title>棲光小島 美術清單</title><style>{css}</style></head><body>{"".join(parts)}</body></html>'
    (D.ROOT / 'docs' / '美術清單.html').write_text(page, encoding='utf-8')
    print('ok', out, 'summary', summary, 'total', tot)


if __name__ == '__main__':
    build()
